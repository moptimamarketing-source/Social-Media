#!/usr/bin/env python3
"""make_calendar — jana kalendar content Excel daripada plan.json.

Guna:  python3 make_calendar.py plan.json kalendar.xlsx
"""
import json
import sys
from collections import Counter
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HARI = ["Isnin", "Selasa", "Rabu", "Khamis", "Jumaat", "Sabtu", "Ahad"]
PILLAR_FILL = {"ilmu": "DCEBFF", "bukti": "DFF5E1", "cerita": "DFF5E1",
               "tabir": "FFF1D6", "jualan": "FFDCDC", "tawaran": "FFDCDC"}
COLS = [("Tarikh", 12), ("Hari", 9), ("Platform", 14), ("Pillar", 12), ("Format", 12),
        ("Tajuk / Hook", 48), ("Pemilik", 12), ("Skill", 16), ("Status", 12), ("Nota", 36)]


def pillar_fill(p):
    p = (p or "").lower()
    for k, c in PILLAR_FILL.items():
        if k in p:
            return PatternFill("solid", fgColor=c)
    return None


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as fh:
        plan = json.load(fh)
    posts = sorted(plan.get("posts", []), key=lambda p: p.get("date", ""))
    if not posts:
        sys.exit("RALAT: tiada 'posts' dalam JSON")

    wb = Workbook()
    ws = wb.active
    ws.title = "Kalendar"
    ws.append([plan.get("title", "Kalendar Content")])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([c for c, _ in COLS])
    for i, (_, w) in enumerate(COLS, 1):
        ws.column_dimensions[chr(64 + i)].width = w
        cell = ws.cell(row=2, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    for p in posts:
        try:
            d = date.fromisoformat(p["date"])
            hari = HARI[d.weekday()]
        except (KeyError, ValueError):
            d, hari = p.get("date", ""), ""
        ws.append([d, hari, p.get("platform", ""), p.get("pillar", ""), p.get("format", ""),
                   p.get("hook", ""), p.get("owner", ""), p.get("skill", ""),
                   p.get("status", "Idea"), p.get("notes", "")])
        r = ws.max_row
        ws.cell(row=r, column=1).number_format = "dd/mm/yyyy"
        f = pillar_fill(p.get("pillar"))
        if f:
            for c in range(1, len(COLS) + 1):
                ws.cell(row=r, column=c).fill = f
        for c in (6, 10):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"Idea,Skrip,Rakam,Edit,Semak,Dijadualkan,Posted"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"I3:I{ws.max_row}")
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:J{ws.max_row}"

    s = wb.create_sheet("Ringkasan")
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 10
    s.column_dimensions["C"].width = 10
    s.append(["Jumlah post", len(posts)])
    s["A1"].font = Font(bold=True)
    for title, key in (("Pillar", "pillar"), ("Platform", "platform"), ("Pemilik", "owner")):
        s.append([])
        s.append([title, "Bil.", "%"])
        s.cell(row=s.max_row, column=1).font = Font(bold=True)
        for k, n in Counter(p.get(key, "") or "(kosong)" for p in posts).most_common():
            s.append([k, n, f"{n / len(posts):.0%}"])
    wb.save(sys.argv[2])
    print(f"SIAP: {sys.argv[2]} ({len(posts)} post)")


if __name__ == "__main__":
    main()
